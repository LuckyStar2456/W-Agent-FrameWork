"""Optional local PDF/Office text templates; dependencies are loaded on use.

These are trusted in-process parsers, not a sandbox for hostile files. No OCR,
formula evaluation, macro execution, network lookup or legacy binary Office.
"""

from __future__ import annotations

import importlib
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile

from w_agent.knowledge_files import Utf8KnowledgeFileParser


class KnowledgeParserDependencyError(RuntimeError):
    """An explicitly selected format requires an unavailable optional package."""


@dataclass(frozen=True, slots=True)
class OptionalDocumentKnowledgeParser:
    """Dispatch UTF-8, PDF, DOCX, XLSX and PPTX; accept at most 1 MiB text.

    Outputs are extraction text, not layout-preserving originals. The host must
    pin dependency versions and change parser_key when extraction behavior changes.
    ZIP limits bound declared expansion, not parser CPU/memory consumption.
    """

    parser_key: str = "optional-documents-v1"
    max_text_bytes: int = 1024 * 1024
    max_zip_bytes: int = 32 * 1024 * 1024
    max_zip_entries: int = 4096
    suffixes: tuple[str, ...] = field(default=(".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx"), init=False)

    def __post_init__(self) -> None:
        from w_agent.knowledge_ingestion import _label

        _label(self.parser_key)
        for value, maximum in ((self.max_text_bytes, 1024 * 1024),
                               (self.max_zip_bytes, 64 * 1024 * 1024), (self.max_zip_entries, 8192)):
            if type(value) is not int or not 1 <= value <= maximum:
                raise ValueError("invalid knowledge parser limit")

    def _dependency(self, module: str):
        try:
            return importlib.import_module(module)
        except ImportError:
            raise KnowledgeParserDependencyError("knowledge-parser-dependency-unavailable:" + module) from None

    def parse(self, data: bytes, *, relative_path: str) -> str:
        suffix = PurePosixPath(relative_path).suffix.lower()
        if suffix not in self.suffixes:
            raise ValueError("unsupported knowledge document format")
        if suffix in (".docx", ".xlsx", ".pptx"):
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if (len(entries) > self.max_zip_entries
                        or sum(item.file_size for item in entries) > self.max_zip_bytes
                        or any(item.flag_bits & 1 for item in entries)):
                    raise ValueError("knowledge Office archive limit exceeded")
        parts: list[str] = []
        size = 0

        def append(text: str) -> None:
            nonlocal size
            if not isinstance(text, str) or "\x00" in text:
                raise ValueError("invalid extracted knowledge text")
            size += len(text.encode("utf-8")) + (1 if parts else 0)
            if size > self.max_text_bytes:
                raise ValueError("knowledge extracted text limit exceeded")
            parts.append(text)

        stream = BytesIO(data)
        if suffix in (".txt", ".md"):
            append(Utf8KnowledgeFileParser().parse(data, relative_path=relative_path))
        elif suffix == ".pdf":
            reader = self._dependency("pypdf").PdfReader(stream)
            if reader.is_encrypted:
                raise ValueError("encrypted knowledge PDF is unsupported")
            for page in reader.pages:
                append(page.extract_text() or "")
        elif suffix == ".docx":
            document = self._dependency("docx").Document(stream)
            # Walk the document body so tables remain in their source order.
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            for element in document.element.body:
                if element.tag.endswith("}p"):
                    append(Paragraph(element, document).text)
                elif element.tag.endswith("}tbl"):
                    for row in Table(element, document).rows:
                        append("\t".join(cell.text for cell in row.cells))
        elif suffix == ".xlsx":
            workbook = self._dependency("openpyxl").load_workbook(stream, read_only=True, data_only=True, keep_links=False)
            try:
                cells = 0
                for sheet in workbook.worksheets:
                    append("[" + sheet.title + "]")
                    for row in sheet.iter_rows(values_only=True):
                        cells += len(row)
                        if cells > 100_000:
                            raise ValueError("knowledge worksheet cell limit exceeded")
                        append("\t".join("" if value is None else str(value) for value in row))
            finally:
                workbook.close()
        else:
            presentation = self._dependency("pptx").Presentation(stream)
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        append(shape.text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
