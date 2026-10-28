"""Actual optional document formats through sync trace entries, not fake parsing."""

import json
from dataclasses import replace

import pytest
import typer
from typer.testing import CliRunner

from w_agent.knowledge_archives import SQLiteKnowledgeRawArchive, archive_knowledge_sources
from w_agent.knowledge_files import KnowledgeFileImportError
from w_agent.knowledge_import_store import SQLiteKnowledgeImportStore
from w_agent.knowledge_sync import KnowledgeSyncError
from w_agent.knowledge_sync_cli import register_knowledge_sync_commands
from w_agent.knowledge_sync_management import KnowledgeSyncImportRequest, import_local_knowledge_sync


def make_document(path, extension):
    dependency = pytest.importorskip({"docx": "docx", "pdf": "pypdf", "xlsx": "openpyxl", "pptx": "pptx"}[extension])
    if extension == "docx":
        document = dependency.Document()
        document.add_paragraph("refund seven days")
        document.save(path)
    elif extension == "xlsx":
        workbook = dependency.Workbook()
        workbook.active.append(["refund seven days"])
        workbook.save(path)
        workbook.close()
    elif extension == "pptx":
        presentation = dependency.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "refund seven days"
        presentation.save(path)
    else:
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = dependency.PdfWriter()
        page = writer.add_blank_page(width=200, height=100)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 10 50 Td (refund seven days) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        writer.write(path)


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx", "pptx"])
async def test_real_format_preview_save_reopen_and_raw_archive(tmp_path, extension):
    path = tmp_path / ("source." + extension)
    make_document(path, extension)
    request = KnowledgeSyncImportRequest(tmp_path / "trace.db", tmp_path, (path.name,), "manual", "v1", confirm_read=True)
    with pytest.raises(KnowledgeFileImportError):
        await import_local_knowledge_sync(request)
    request = replace(request, document_formats=True)
    preview = await import_local_knowledge_sync(request)
    assert not request.trace_database.exists()
    saved = await import_local_knowledge_sync(replace(request, confirm_record=True, confirm_create=True,
        expected_preview_digest=preview["preview_digest"]))
    trace = await SQLiteKnowledgeImportStore(request.trace_database).load("manual", "v1")
    assert trace.digest == saved["trace_digest"]
    assert "refund seven days" in trace.snapshot.documents[0].text
    archive = SQLiteKnowledgeRawArchive(tmp_path / "raw.db")
    await archive.initialize(confirm_create=True)
    receipt = await archive_knowledge_sources(trace, tmp_path, archive, expected_trace_digest=trace.digest,
                                             confirm_read=True, confirm_write=True)
    assert await archive.load(receipt.blob_digests[0], confirm_read=True) == path.read_bytes()


@pytest.mark.asyncio
async def test_parser_selection_and_limits_in_preview_lock(tmp_path, monkeypatch):
    from w_agent.knowledge_parsers import OptionalDocumentKnowledgeParser
    import w_agent.knowledge_sync_management as management
    (tmp_path / "source.md").write_text("same source", encoding="utf-8")
    request = KnowledgeSyncImportRequest(tmp_path / "trace.db", tmp_path, ("source.md",), "manual", "v1", confirm_read=True)
    utf8 = await import_local_knowledge_sync(request)
    optional = await import_local_knowledge_sync(replace(request, document_formats=True))
    assert utf8["preview_digest"] != optional["preview_digest"]
    with pytest.raises(KnowledgeSyncError, match="preview-changed"):
        await import_local_knowledge_sync(replace(request, document_formats=True, confirm_record=True, confirm_create=True,
            expected_preview_digest=utf8["preview_digest"]))
    monkeypatch.setattr(management, "OptionalDocumentKnowledgeParser", lambda: OptionalDocumentKnowledgeParser(max_zip_entries=4095))
    with pytest.raises(KnowledgeSyncError, match="preview-changed"):
        await import_local_knowledge_sync(replace(request, document_formats=True, confirm_record=True, confirm_create=True,
            expected_preview_digest=optional["preview_digest"]))
    assert not request.trace_database.exists()


@pytest.mark.asyncio
async def test_document_source_drift_prevents_save(tmp_path):
    path = tmp_path / "source.docx"
    make_document(path, "docx")
    request = KnowledgeSyncImportRequest(tmp_path / "trace.db", tmp_path, (path.name,), "manual", "v1", confirm_read=True, document_formats=True)
    preview = await import_local_knowledge_sync(request)
    import docx
    document = docx.Document(path)
    document.add_paragraph("changed")
    document.save(path)
    with pytest.raises(KnowledgeSyncError, match="preview-changed"):
        await import_local_knowledge_sync(replace(request, confirm_record=True, confirm_create=True,
            expected_preview_digest=preview["preview_digest"]))
    assert not request.trace_database.exists()


def test_cli_docx_explicit_enable_and_locked_save(tmp_path):
    make_document(tmp_path / "source.docx", "docx")
    app = typer.Typer()
    register_knowledge_sync_commands(app)
    args = ["knowledge-sync-import", "--trace-database", str(tmp_path / "trace.db"), "--root", str(tmp_path),
            "--path", "source.docx", "--namespace", "manual", "--version", "v1", "--confirm-read"]
    runner = CliRunner()
    assert runner.invoke(app, args).exit_code == 2
    args.append("--document-formats")
    preview = runner.invoke(app, args)
    assert preview.exit_code == 0, preview.output
    result = runner.invoke(app, args + ["--confirm-record", "--confirm-create", "--expected-preview-digest", json.loads(preview.output)["preview_digest"]])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["recorded"]
